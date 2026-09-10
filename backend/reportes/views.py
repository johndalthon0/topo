from datetime import datetime, date
from decimal import Decimal
import io
import os

from django.conf import settings
from django.http import HttpResponse
from django.utils import timezone
from django.db.models import Sum

from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status

from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from .models import Reporte
from compras.models import Compra, DetalleCompra
from ventas.models import Venta, DetalleVenta
from medicamentos.models import Medicamento


# ============================================================
# FUNCIONES AUXILIARES
# ============================================================

def es_anulado(estado):
    return str(estado or '').strip().upper() == 'ANULADA'


def obtener_fecha(valor):
    if not valor:
        return None

    if isinstance(valor, date):
        return valor

    if isinstance(valor, str):
        try:
            return datetime.strptime(valor, '%Y-%m-%d').date()
        except ValueError:
            return None

    return None


def filtrar_por_fecha(queryset, fecha_inicio, fecha_fin):
    if fecha_inicio:
        queryset = queryset.filter(fecha__gte=fecha_inicio)

    if fecha_fin:
        queryset = queryset.filter(fecha__lte=fecha_fin)

    return queryset


def nombre_usuario_venta(venta):
    if venta.usuario:
        return getattr(venta.usuario, 'nombre', None) or str(venta.usuario)

    return 'Administrador'


def nombre_proveedor(compra):
    if compra.proveedor:
        return getattr(compra.proveedor, 'nombre', None) or str(compra.proveedor)

    return 'Sin proveedor'


def fecha_registro_medicamento(medicamento):
    for campo in [
        'fecha_registro',
        'fecha_creacion',
        'created_at',
        'fecha'
    ]:
        valor = getattr(medicamento, campo, None)
        if valor:
            return str(valor)[:10]

    return ''


def precio_compra_medicamento(medicamento):
    return Decimal(str(getattr(medicamento, 'precio_compra', 0) or 0))


def stock_medicamento(medicamento):
    return int(getattr(medicamento, 'stock', 0) or 0)


def aplicar_tipo_reporte(tipo_reporte):
    texto = str(tipo_reporte or 'General').strip().lower()

    return {
        'general': 'General',
        'compras': 'Compras',
        'ventas': 'Ventas',
        'medicamentos': 'Medicamentos',
        'inventario': 'Medicamentos',
    }.get(texto, 'General')


# ============================================================
# RESUMEN GENERAL
# ============================================================

class ResumenReportesView(APIView):

    def get(self, request):

        compras_validas = Compra.objects.exclude(
            estado__iexact='ANULADA'
        )

        ventas_validas = Venta.objects.exclude(
            estado__iexact='ANULADA'
        )

        total_compras = compras_validas.aggregate(
            total=Sum('total')
        )['total'] or Decimal('0.00')

        total_ventas = ventas_validas.aggregate(
            total=Sum('total')
        )['total'] or Decimal('0.00')

        ganancia = total_ventas - total_compras

        valor_inventario = Decimal('0.00')
        medicamentos = Medicamento.objects.all()

        for medicamento in medicamentos:
            valor_inventario += (
                Decimal(stock_medicamento(medicamento))
                * precio_compra_medicamento(medicamento)
            )

        return Response({
            'totalCompras': float(total_compras),
            'totalVentas': float(total_ventas),
            'ganancia': float(ganancia),
            'valorInventario': float(valor_inventario),
            'totalMedicamentos': medicamentos.count(),
            'cantidadComprasValidas': compras_validas.count(),
            'cantidadVentasValidas': ventas_validas.count(),
            'comprasAnuladas': Compra.objects.filter(
                estado__iexact='ANULADA'
            ).count(),
            'ventasAnuladas': Venta.objects.filter(
                estado__iexact='ANULADA'
            ).count(),
        }, status=status.HTTP_200_OK)


# ============================================================
# COMPRAS PARA REPORTES
# ============================================================

class ComprasReportesView(APIView):

    def get(self, request):

        compras = (
            Compra.objects
            .select_related('proveedor', 'compra_origen')
            .all()
            .order_by('-fecha', '-id')
        )

        datos = []

        for compra in compras:

            detalle = []

            for item in (
                DetalleCompra.objects
                .filter(compra=compra)
                .select_related('medicamento')
            ):
                detalle.append({
                    'medicamento': item.medicamento.nombre,
                    'medicamento_id': item.medicamento.id,
                    'cantidad': item.cantidad,
                    'precio_compra': float(item.precio_compra),
                    'subtotal': float(item.subtotal),
                })

            datos.append({
                'id': compra.id,
                'factura': compra.factura,
                'proveedor': nombre_proveedor(compra),
                'fecha': str(compra.fecha),
                'total': float(compra.total),
                'estado': compra.estado,
                'observacion': compra.observacion or '',
                'usuario': compra.anulado_por or 'Administrador',
                'motivo_anulacion': compra.motivo_anulacion,
                'fecha_anulacion': (
                    compra.fecha_anulacion.isoformat()
                    if compra.fecha_anulacion
                    else None
                ),
                'anulado_por': compra.anulado_por,
                'compra_origen_id': compra.compra_origen_id,
                'compra_origen_factura': (
                    compra.compra_origen.factura
                    if compra.compra_origen
                    else None
                ),
                'tipo_registro': (
                    'CORRECCION_COMPRA'
                    if compra.compra_origen_id
                    else 'COMPRA'
                ),
                'hash_operacion': compra.hash_operacion,
                'transaction_hash': compra.transaction_hash,
                'block_number': compra.block_number,
                'detalle': detalle,
            })

        return Response(datos, status=status.HTTP_200_OK)


# ============================================================
# VENTAS PARA REPORTES
# ============================================================

class VentasReportesView(APIView):

    def get(self, request):

        ventas = (
            Venta.objects
            .select_related('usuario', 'venta_origen')
            .all()
            .order_by('-fecha', '-id')
        )

        datos = []

        for venta in ventas:

            detalle = []

            for item in (
                DetalleVenta.objects
                .filter(venta=venta)
                .select_related('medicamento')
            ):
                detalle.append({
                    'medicamento': item.medicamento.nombre,
                    'medicamento_id': item.medicamento.id,
                    'cantidad': item.cantidad,
                    'precio_venta': float(item.precio_venta),
                    'subtotal': float(item.subtotal),
                })

            datos.append({
                'id': venta.id,
                'numero_venta': venta.numero_venta,
                'cliente': getattr(venta, 'cliente', '') or 'Consumidor Final',
                'fecha': str(venta.fecha),
                'total': float(venta.total),
                'estado': venta.estado,
                'observacion': venta.observacion or '',
                'usuario': nombre_usuario_venta(venta),
                'motivo_anulacion': venta.motivo_anulacion,
                'fecha_anulacion': (
                    venta.fecha_anulacion.isoformat()
                    if venta.fecha_anulacion
                    else None
                ),
                'anulado_por': venta.anulado_por,
                'venta_origen_id': venta.venta_origen_id,
                'venta_origen_numero': (
                    venta.venta_origen.numero_venta
                    if venta.venta_origen
                    else None
                ),
                'tipo_registro': (
                    'CORRECCION_VENTA'
                    if venta.venta_origen_id
                    else 'VENTA'
                ),
                'hash_operacion': venta.hash_operacion,
                'transaction_hash': venta.transaction_hash,
                'block_number': venta.block_number,
                'detalle': detalle,
            })

        return Response(datos, status=status.HTTP_200_OK)


# ============================================================
# MEDICAMENTOS PARA REPORTES
# ============================================================

class MedicamentosReportesView(APIView):

    def get(self, request):

        medicamentos = Medicamento.objects.all().order_by('nombre')
        datos = []

        for medicamento in medicamentos:

            precio = precio_compra_medicamento(medicamento)
            stock = stock_medicamento(medicamento)

            datos.append({
                'id': medicamento.id,
                'nombre': medicamento.nombre,
                'stock': stock,
                'precio_compra': float(precio),
                'valor_inventario': float(
                    Decimal(stock) * precio
                ),
                'fecha_registro': fecha_registro_medicamento(medicamento),
                'estado': 'Correcto',
            })

        return Response(datos, status=status.HTTP_200_OK)


# ============================================================
# EXPORTAR PDF
# ============================================================

class ExportarPDFView(APIView):

    def post(self, request):

        fecha_inicio = obtener_fecha(
            request.data.get('fecha_inicio')
        )

        fecha_fin = obtener_fecha(
            request.data.get('fecha_fin')
        )

        tipo_reporte = aplicar_tipo_reporte(
            request.data.get('tipo_reporte')
        )

        compras = filtrar_por_fecha(
            Compra.objects.all().order_by('-fecha', '-id'),
            fecha_inicio,
            fecha_fin
        )

        ventas = filtrar_por_fecha(
            Venta.objects.all().order_by('-fecha', '-id'),
            fecha_inicio,
            fecha_fin
        )

        compras_validas = compras.exclude(
            estado__iexact='ANULADA'
        )

        ventas_validas = ventas.exclude(
            estado__iexact='ANULADA'
        )

        total_compras = compras_validas.aggregate(
            total=Sum('total')
        )['total'] or Decimal('0.00')

        total_ventas = ventas_validas.aggregate(
            total=Sum('total')
        )['total'] or Decimal('0.00')

        ganancia = total_ventas - total_compras

        medicamentos = Medicamento.objects.all().order_by('nombre')

        valor_inventario = Decimal('0.00')

        for medicamento in medicamentos:
            valor_inventario += (
                Decimal(stock_medicamento(medicamento))
                * precio_compra_medicamento(medicamento)
            )

        carpeta_reportes = os.path.join(
            settings.MEDIA_ROOT,
            'reportes'
        )

        os.makedirs(
            carpeta_reportes,
            exist_ok=True
        )

        inicio_nombre = str(fecha_inicio) if fecha_inicio else 'todas'
        fin_nombre = str(fecha_fin) if fecha_fin else 'todas'

        nombre_archivo = (
            f'reporte_{tipo_reporte.lower()}_'
            f'{inicio_nombre}_{fin_nombre}.pdf'
        )

        ruta_archivo = os.path.join(
            carpeta_reportes,
            nombre_archivo
        )

        pdf = canvas.Canvas(
            ruta_archivo,
            pagesize=letter
        )

        ancho, alto = letter
        y = alto - 45

        def nueva_pagina():
            nonlocal y
            pdf.showPage()
            y = alto - 45

        def linea(texto, tamano=9, negrita=False, salto=16):
            nonlocal y

            if y < 55:
                nueva_pagina()

            pdf.setFont(
                'Helvetica-Bold' if negrita else 'Helvetica',
                tamano
            )

            pdf.drawString(45, y, str(texto)[:110])
            y -= salto

        pdf.setTitle(
            'Reporte - Farmacia Coria'
        )

        linea('FARMACIA CORIA', 15, True, 22)
        linea(f'REPORTE: {tipo_reporte.upper()}', 12, True, 20)
        linea(
            'Periodo: '
            + (str(fecha_inicio) if fecha_inicio else 'Sin límite')
            + ' hasta '
            + (str(fecha_fin) if fecha_fin else 'Sin límite'),
            9,
            False,
            22
        )

        linea('RESUMEN', 11, True, 18)
        linea(f'Total compras válidas: {total_compras:.2f} Bs')
        linea(f'Total ventas válidas: {total_ventas:.2f} Bs')
        linea(f'Ganancia: {ganancia:.2f} Bs')
        linea(f'Valor actual del inventario: {valor_inventario:.2f} Bs')
        linea('Las operaciones ANULADAS no se incluyen en los totales.', 8, True, 22)

        if tipo_reporte in ['General', 'Compras']:
            linea('COMPRAS', 11, True, 18)

            for compra in compras:
                texto = (
                    f'#{compra.id} | {compra.fecha} | '
                    f'{compra.factura} | {compra.total:.2f} Bs | '
                    f'{compra.estado}'
                )
                linea(texto)

                if es_anulado(compra.estado) and compra.motivo_anulacion:
                    linea(
                        f'   Motivo: {compra.motivo_anulacion}',
                        8,
                        False,
                        13
                    )

        if tipo_reporte in ['General', 'Ventas']:
            linea('VENTAS', 11, True, 18)

            for venta in ventas:
                texto = (
                    f'#{venta.id} | {venta.fecha} | '
                    f'{venta.numero_venta} | {venta.total:.2f} Bs | '
                    f'{venta.estado}'
                )
                linea(texto)

                if es_anulado(venta.estado) and venta.motivo_anulacion:
                    linea(
                        f'   Motivo: {venta.motivo_anulacion}',
                        8,
                        False,
                        13
                    )

        if tipo_reporte in ['General', 'Medicamentos']:
            linea('MEDICAMENTOS / INVENTARIO', 11, True, 18)

            for medicamento in medicamentos:
                precio = precio_compra_medicamento(medicamento)
                stock = stock_medicamento(medicamento)
                valor = Decimal(stock) * precio

                linea(
                    f'#{medicamento.id} | {medicamento.nombre} | '
                    f'Stock: {stock} | Valor: {valor:.2f} Bs'
                )

        linea(
            'Reporte generado automáticamente por el sistema.',
            8,
            False,
            12
        )

        pdf.save()

        fecha_reporte_inicio = fecha_inicio or timezone.now().date()
        fecha_reporte_fin = fecha_fin or timezone.now().date()

        Reporte.objects.create(
            tipo_reporte=tipo_reporte,
            fecha_inicio=fecha_reporte_inicio,
            fecha_fin=fecha_reporte_fin,
            archivo=f'reportes/{nombre_archivo}'
        )

        with open(ruta_archivo, 'rb') as archivo:
            respuesta = HttpResponse(
                archivo.read(),
                content_type='application/pdf'
            )

        respuesta['Content-Disposition'] = (
            f'attachment; filename="{nombre_archivo}"'
        )

        return respuesta


# ============================================================
# EXPORTAR EXCEL
# ============================================================

class ExportarExcelView(APIView):

    def post(self, request):

        fecha_inicio = obtener_fecha(
            request.data.get('fecha_inicio')
        )

        fecha_fin = obtener_fecha(
            request.data.get('fecha_fin')
        )

        tipo_reporte = aplicar_tipo_reporte(
            request.data.get('tipo_reporte')
        )

        compras = filtrar_por_fecha(
            Compra.objects.all().order_by('-fecha', '-id'),
            fecha_inicio,
            fecha_fin
        )

        ventas = filtrar_por_fecha(
            Venta.objects.all().order_by('-fecha', '-id'),
            fecha_inicio,
            fecha_fin
        )

        medicamentos = Medicamento.objects.all().order_by('nombre')

        compras_validas = compras.exclude(
            estado__iexact='ANULADA'
        )

        ventas_validas = ventas.exclude(
            estado__iexact='ANULADA'
        )

        total_compras = sum(
            (compra.total for compra in compras_validas),
            Decimal('0.00')
        )

        total_ventas = sum(
            (venta.total for venta in ventas_validas),
            Decimal('0.00')
        )

        ganancia = total_ventas - total_compras

        valor_inventario = Decimal('0.00')

        for medicamento in medicamentos:
            valor_inventario += (
                Decimal(stock_medicamento(medicamento))
                * precio_compra_medicamento(medicamento)
            )

        libro = Workbook()
        hoja_resumen = libro.active
        hoja_resumen.title = 'Resumen'

        hoja_resumen['A1'] = 'FARMACIA CORIA'
        hoja_resumen['A1'].font = Font(
            bold=True,
            size=16
        )

        hoja_resumen['A2'] = f'REPORTE {tipo_reporte.upper()}'
        hoja_resumen['A4'] = 'Fecha inicio'
        hoja_resumen['B4'] = fecha_inicio or 'Sin límite'
        hoja_resumen['A5'] = 'Fecha fin'
        hoja_resumen['B5'] = fecha_fin or 'Sin límite'
        hoja_resumen['A7'] = 'Total compras válidas'
        hoja_resumen['B7'] = float(total_compras)
        hoja_resumen['A8'] = 'Total ventas válidas'
        hoja_resumen['B8'] = float(total_ventas)
        hoja_resumen['A9'] = 'Ganancia'
        hoja_resumen['B9'] = float(ganancia)
        hoja_resumen['A10'] = 'Valor inventario'
        hoja_resumen['B10'] = float(valor_inventario)
        hoja_resumen['A12'] = 'Nota'
        hoja_resumen['B12'] = (
            'Las operaciones ANULADAS se conservan en el historial, '
            'pero no se incluyen en los totales.'
        )

        if tipo_reporte in ['General', 'Compras']:
            hoja_compras = libro.create_sheet('Compras')

            encabezados_compras = [
                'ID',
                'Fecha',
                'Factura',
                'Proveedor',
                'Total',
                'Estado',
                'Tipo',
                'Referencia original',
                'Motivo anulación',
                'Anulado por'
            ]

            hoja_compras.append(encabezados_compras)

            for celda in hoja_compras[1]:
                celda.font = Font(bold=True)
                celda.alignment = Alignment(horizontal='center')

            for compra in compras:
                hoja_compras.append([
                    compra.id,
                    compra.fecha,
                    compra.factura,
                    nombre_proveedor(compra),
                    float(compra.total),
                    compra.estado,
                    (
                        'Compra corregida'
                        if compra.compra_origen_id
                        else 'Compra original'
                    ),
                    (
                        compra.compra_origen.factura
                        if compra.compra_origen
                        else ''
                    ),
                    compra.motivo_anulacion or '',
                    compra.anulado_por or '',
                ])

        if tipo_reporte in ['General', 'Ventas']:
            hoja_ventas = libro.create_sheet('Ventas')

            encabezados_ventas = [
                'ID',
                'Fecha',
                'Número venta',
                'Cliente',
                'Usuario',
                'Total',
                'Estado',
                'Tipo',
                'Referencia original',
                'Motivo anulación',
                'Anulado por'
            ]

            hoja_ventas.append(encabezados_ventas)

            for celda in hoja_ventas[1]:
                celda.font = Font(bold=True)
                celda.alignment = Alignment(horizontal='center')

            for venta in ventas:
                hoja_ventas.append([
                    venta.id,
                    venta.fecha,
                    venta.numero_venta,
                    getattr(venta, 'cliente', '') or 'Consumidor Final',
                    nombre_usuario_venta(venta),
                    float(venta.total),
                    venta.estado,
                    (
                        'Venta corregida'
                        if venta.venta_origen_id
                        else 'Venta original'
                    ),
                    (
                        venta.venta_origen.numero_venta
                        if venta.venta_origen
                        else ''
                    ),
                    venta.motivo_anulacion or '',
                    venta.anulado_por or '',
                ])

        if tipo_reporte in ['General', 'Medicamentos']:
            hoja_medicamentos = libro.create_sheet('Medicamentos')

            encabezados_medicamentos = [
                'ID',
                'Nombre',
                'Stock',
                'Precio compra',
                'Valor inventario'
            ]

            hoja_medicamentos.append(encabezados_medicamentos)

            for celda in hoja_medicamentos[1]:
                celda.font = Font(bold=True)
                celda.alignment = Alignment(horizontal='center')

            for medicamento in medicamentos:
                precio = precio_compra_medicamento(medicamento)
                stock = stock_medicamento(medicamento)

                hoja_medicamentos.append([
                    medicamento.id,
                    medicamento.nombre,
                    stock,
                    float(precio),
                    float(Decimal(stock) * precio),
                ])

        for hoja in libro.worksheets:
            for columna in hoja.columns:
                longitud = 0
                letra = columna[0].column_letter

                for celda in columna:
                    if celda.value is not None:
                        longitud = max(
                            longitud,
                            len(str(celda.value))
                        )

                hoja.column_dimensions[letra].width = min(
                    longitud + 3,
                    45
                )

        archivo_memoria = io.BytesIO()
        libro.save(archivo_memoria)
        archivo_memoria.seek(0)

        inicio_nombre = str(fecha_inicio) if fecha_inicio else 'todas'
        fin_nombre = str(fecha_fin) if fecha_fin else 'todas'

        nombre_archivo = (
            f'reporte_{tipo_reporte.lower()}_'
            f'{inicio_nombre}_{fin_nombre}.xlsx'
        )

        fecha_reporte_inicio = fecha_inicio or timezone.now().date()
        fecha_reporte_fin = fecha_fin or timezone.now().date()

        Reporte.objects.create(
            tipo_reporte=tipo_reporte,
            fecha_inicio=fecha_reporte_inicio,
            fecha_fin=fecha_reporte_fin,
            archivo=f'reportes/{nombre_archivo}'
        )

        respuesta = HttpResponse(
            archivo_memoria.getvalue(),
            content_type=(
                'application/vnd.openxmlformats-officedocument.'
                'spreadsheetml.sheet'
            )
        )

        respuesta['Content-Disposition'] = (
            f'attachment; filename="{nombre_archivo}"'
        )

        return respuesta